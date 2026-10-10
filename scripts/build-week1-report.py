from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "dist" / "Bao_cao_tuan_1_Learner_Lab_Ops_Dashboard_24110019.docx"
SCHOOL_LOGO = ROOT / "assets" / "hcmute-logo-from-reference.png"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)

def borders(cell, color="D9D9D9"):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.first_child_found_in("w:tcBorders")
    if b is None:
        b = OxmlElement("w:tcBorders")
        tcPr.append(b)
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), color)
        b.append(node)

def set_cell_text(cell, value, bold=False, size=9.5):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(str(value))
    r.bold = bold
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    t.autofit = False
    for i, value in enumerate(headers):
        c = t.rows[0].cells[i]
        set_cell_text(c, value, True)
        shade(c, "ECECEC")
        borders(c)
        if widths: c.width = Cm(widths[i])
    for ri, vals in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(vals):
            set_cell_text(cells[i], value)
            borders(cells[i])
            if widths: cells[i].width = Cm(widths[i])
            if ri % 2: shade(cells[i], "F8F8F8")
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def para(text="", style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        p.add_run(bold_prefix).bold = True
        p.add_run(text[len(bold_prefix):])
    else:
        p.add_run(text)
    return p

def hyperlink(paragraph, text, url):
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(color)
    props.append(underline)
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    link.append(run)
    paragraph._p.append(link)
    return link

def heading(text, level=1):
    doc.add_heading(text, level=level)

def caption(text):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)

def picture(path, width_cm, label):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    caption(label)

def set_page_border(section):
    sectPr = section._sectPr
    pgBorders = OxmlElement("w:pgBorders")
    pgBorders.set(qn("w:offsetFrom"), "page")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "8")
        e.set(qn("w:space"), "24")
        e.set(qn("w:color"), "333333")
        pgBorders.append(e)
    sectPr.append(pgBorders)

def page_number(paragraph):
    paragraph.add_run("Trang ")
    r = paragraph.add_run()
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    r._r.addnext(fld)

def insert_toc():
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), r'TOC \o "1-2" \h \z \u')
    p._p.append(field)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin = Cm(2.3)
sec.bottom_margin = Cm(2.0)
sec.left_margin = sec.right_margin = Cm(2.35)
set_page_border(sec)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.22
normal.paragraph_format.space_after = Pt(6)
for st_name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 12)):
    st = styles[st_name]
    st.font.name = "Times New Roman"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(13 if st_name != "Title" else 0)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.keep_with_next = True
title_ppr = styles["Title"]._element.get_or_add_pPr()
for title_border in title_ppr.findall(qn("w:pBdr")):
    title_ppr.remove(title_border)
styles["Caption"].font.name = "Times New Roman"
styles["Caption"].font.size = Pt(9)
styles["Caption"].font.italic = True
styles["Caption"].font.color.rgb = RGBColor(0, 0, 0)

# Trang bìa lấy nhịp bố cục từ báo cáo tham khảo, chỉ dùng dữ kiện của đề tài này.
for text, size, space in [
    ("TRƯỜNG ĐẠI HỌC CÔNG NGHỆ KỸ THUẬT TP HCM", 14, 5),
    ("KHOA CÔNG NGHỆ THÔNG TIN", 13, 10),
    ("MÔN HỌC ĐIỆN TOÁN ĐÁM MÂY", 15, 6),
    ("CLOUD - ĐỢT 1 - 2026–2027", 11, 9),
    ("BÁO CÁO TIẾN ĐỘ TUẦN 1", 18, 12),
    ("BẢNG ĐIỀU KHIỂN VẬN HÀNH", 17, 3),
    ("CHO AWS LEARNER LAB", 17, 42),
]:
    p = doc.add_paragraph()
    if text in ("BẢNG ĐIỀU KHIỂN VẬN HÀNH", "CHO AWS LEARNER LAB"):
        p.style = styles["Title"]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space)
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(size)
    if text == "CLOUD - ĐỢT 1 - 2026–2027":
        logo_p = doc.add_paragraph()
        logo_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo_p.paragraph_format.space_after = Pt(18)
        logo_p.add_run().add_picture(str(SCHOOL_LOGO), width=Cm(3.8))

cover = doc.add_table(rows=4, cols=1)
cover.autofit = False
cover.alignment = WD_TABLE_ALIGNMENT.CENTER
for c, text in zip([r.cells[0] for r in cover.rows], [
    "Nhóm lĩnh vực: Monitoring & Operations",
    "Giảng viên hướng dẫn: Huỳnh Xuân Phụng",
    "Sinh viên: Lê Minh Huy — MSSV 24110019",
    "MÃ HỌC PHẦN: CLCO332779",
]):
    set_cell_text(c, text, bold=text.startswith("MÃ HỌC PHẦN:"), size=10.5)
    borders(c, "666666")
doc.add_paragraph("\n\n\n")
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("Ngày 10 tháng 10 năm 2026")
r.italic = True
r.font.size = Pt(11)

bodysec = doc.add_section(WD_SECTION.NEW_PAGE)
bodysec.page_width, bodysec.page_height = Cm(21), Cm(29.7)
bodysec.top_margin = Cm(2.1)
bodysec.bottom_margin = Cm(2.0)
bodysec.left_margin = bodysec.right_margin = Cm(2.3)
for inherited_border in bodysec._sectPr.findall(qn("w:pgBorders")):
    bodysec._sectPr.remove(inherited_border)
bodysec.footer.is_linked_to_previous = False
fp = bodysec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
page_number(fp)

toc_heading = doc.add_paragraph(style="Title")
toc_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
toc_heading.add_run("MỤC LỤC")
insert_toc()
doc.add_page_break()

heading("TÓM TẮT", 1)
para("Trong tuần 1, đề tài tập trung tìm hiểu các dịch vụ cần dùng, xây dựng kiến trúc và sơ đồ, chuẩn bị VPC, kiểm tra LabRole có sẵn, tạo bucket S3 và gắn tag cho tài nguyên. Phần thu thập dữ liệu tự động, API, giao diện biểu đồ và cảnh báo được dành cho các tuần sau.")
para("Nền tảng được triển khai trên tài khoản Learner Lab 213858341631 ở Region us-east-1. Ngày 08/10/2026, stack learnerlab-ops-week1 hoàn tất với 10/10 tài nguyên ở trạng thái CREATE_COMPLETE. Kết quả gồm VPC 10.20.0.0/16, hai subnet ở hai Availability Zone, tuyến qua Internet Gateway và một bucket S3 riêng tư, mã hóa SSE-S3. LabRole đã được kiểm tra bằng IAM GetRole; LabInstanceProfile chưa được kiểm tra riêng.")

heading("Chương 1 Tổng quan và phạm vi tuần 1", 1)
heading("1.1 Bài toán và đầu ra", 2)
para("Đề tài hướng tới một trang web giúp theo dõi tài nguyên đang chạy, chi phí AWS ghi nhận và cảnh báo khi vượt ngưỡng ngân sách Learner Lab. Theo thiết kế, Lambda sẽ đọc thông tin EC2, RDS, SageMaker và EKS; DynamoDB lưu bản ghi theo thời điểm; API cung cấp dữ liệu cho trang web; SNS gửi cảnh báo. Tuần 1 tập trung vào thiết kế và hạ tầng ban đầu.")
table(["Đầu việc tuần 1", "Sản phẩm hiện có", "Trạng thái"], [
    ("Tìm hiểu dịch vụ và giới hạn Lab", "docs/week-1/01-services-and-requirements.md", "Đã có tài liệu"),
    ("Thiết kế kiến trúc, sơ đồ", "docs/week-1/02-architecture.md; docs/diagrams/architecture-week1-simple.png", "Đã có bản thiết kế"),
    ("Dựng VPC, kiểm tra LabRole, tạo S3", "infra/week1-foundation.yaml; stack AWS", "Đã triển khai 08/10"),
    ("Gắn tag và lưu ảnh", "evidence/week-1/02–04", "Đã kiểm tra 08/10"),
    ("Nhật ký công việc", "docs/week-1/worklog.csv", "Đã ghi việc; chưa có số giờ thực tế"),
], [4.1, 8.2, 4.0])
caption("Bảng 1.1. Công việc và sản phẩm của tuần 1")

heading("1.2 Dịch vụ đã tìm hiểu", 2)
table(["Dịch vụ", "Vai trò trong kiến trúc", "Tình trạng tuần 1"], [
    ("VPC, IAM LabRole, S3", "Mạng, quyền truy cập và nơi lưu tệp dự án", "VPC/S3 đã triển khai; LabRole đã kiểm tra"),
    ("Lambda, EventBridge", "Kiểm kê tài nguyên theo lịch", "Mới ở bước thiết kế"),
    ("DynamoDB, API Gateway", "Lưu bản ghi và cung cấp dữ liệu cho web", "Mới ở bước thiết kế"),
    ("Cost Explorer, SNS", "Đọc chi phí và gửi cảnh báo", "Mới ở bước thiết kế"),
    ("CloudWatch", "Theo dõi log và lỗi của Lambda", "Dự kiến triển khai sau"),
], [3.7, 7.6, 5.0])
caption("Bảng 1.2. Vai trò và tiến độ của các dịch vụ AWS")

heading("Chương 2 Thiết kế kiến trúc", 1)
heading("2.1 Luồng dự kiến", 2)
para("Kiến trúc dự kiến dùng hai lịch EventBridge: một lịch kiểm kê tài nguyên, một lịch lấy thông tin chi phí. Lambda kiểm kê đọc dữ liệu tại us-east-1 và us-west-2. Lambda chi phí gọi Cost Explorer với tần suất giới hạn, sau đó lưu kết quả vào DynamoDB. Trang web đọc dữ liệu đã lưu qua API để tránh gọi Cost Explorer mỗi lần mở trang. Khi chi phí vượt ngưỡng, SNS sẽ gửi cảnh báo. S3 lưu tệp giao diện và các tệp của dự án. Luồng này mới được thiết kế, chưa vận hành trong tuần 1.")
heading("2.2 Mã nguồn và tài liệu tuần 1", 2)
para("Mã nguồn, template và minh chứng của đề tài được lưu trong repository GitHub công khai. Các tài liệu phân biệt rõ phần đã triển khai trên Learner Lab với các thành phần mới dừng ở bước thiết kế.")
repo_link = doc.add_paragraph()
repo_link.paragraph_format.left_indent = Cm(0.3)
repo_link.paragraph_format.space_after = Pt(5)
hyperlink(repo_link, "github.com/leminhhuy-ute/learnerlab-ops-dashboard",
          "https://github.com/leminhhuy-ute/learnerlab-ops-dashboard")
table(["Tệp / thư mục", "Nội dung và trạng thái tuần 1"], [
    ("infra/week1-foundation.yaml", "Template VPC, subnet, Internet Gateway, route và bucket S3 đã triển khai."),
    ("docs/week-1/", "Kế hoạch, tìm hiểu dịch vụ, kiến trúc, runbook, kiểm tra và biên bản AWS."),
    ("docs/diagrams/architecture-week1-simple.png", "Sơ đồ một trang; nét liền là nền tảng đã kiểm tra, nét đứt là luồng dự kiến."),
    ("evidence/week-1/", "Ảnh Learner Lab, CloudFormation, cấu hình mạng và bảo mật S3."),
    ("docs/week-1/worklog.csv", "Đầu việc và mốc AWS có nguồn; giờ cá nhân để trống khi chưa có dữ liệu xác nhận."),
], [6.1, 10.2])
caption("Bảng 2.1. Tệp dự án và minh chứng của tuần 1")
para("Stack tuần 1 chỉ tạo mạng và bucket riêng tư. Lambda, EventBridge, DynamoDB, API Gateway, Cost Explorer và SNS trong sơ đồ là các phần dự kiến cho giai đoạn tiếp theo, chưa được báo cáo là đã chạy.")
diagram_section = doc.add_section(WD_SECTION.NEW_PAGE)
diagram_section.orientation = WD_ORIENT.LANDSCAPE
diagram_section.page_width, diagram_section.page_height = Cm(29.7), Cm(21)
diagram_section.top_margin = diagram_section.bottom_margin = Cm(1.2)
diagram_section.left_margin = diagram_section.right_margin = Cm(1.3)
for inherited_border in diagram_section._sectPr.findall(qn("w:pgBorders")):
    diagram_section._sectPr.remove(inherited_border)
diagram_section.footer.is_linked_to_previous = False
page_number(diagram_section.footer.paragraphs[0])
picture(ROOT / "docs" / "diagrams" / "architecture-week1-simple.png", 26.6,
        "Hình 2.1. Nền tảng tuần 1 và luồng kiến trúc dự kiến")

body_after_diagram = doc.add_section(WD_SECTION.NEW_PAGE)
body_after_diagram.orientation = WD_ORIENT.PORTRAIT
body_after_diagram.page_width, body_after_diagram.page_height = Cm(21), Cm(29.7)
body_after_diagram.top_margin = Cm(2.1)
body_after_diagram.bottom_margin = Cm(2.0)
body_after_diagram.left_margin = body_after_diagram.right_margin = Cm(2.3)
for inherited_border in body_after_diagram._sectPr.findall(qn("w:pgBorders")):
    body_after_diagram._sectPr.remove(inherited_border)
body_after_diagram.footer.is_linked_to_previous = False
page_number(body_after_diagram.footer.paragraphs[0])
heading("2.3 Ranh giới đã triển khai", 2)
para("Stack tuần 1 tạo VPC, hai subnet và bucket S3. Hai subnet có tuyến 0.0.0.0/0 qua Internet Gateway, nhưng MapPublicIpOnLaunch=False nên tài nguyên đặt trong subnet không tự nhận public IPv4. Stack không tạo EC2, NAT Gateway, Elastic IP, RDS, EKS, SageMaker, Lambda, DynamoDB hoặc SNS. LabRole là vai trò có sẵn trong Learner Lab, không phải tài nguyên do template tạo.")

heading("Chương 3 Kết quả triển khai trên AWS Learner Lab", 1)
heading("3.1 Tài khoản, thời điểm và CloudFormation", 2)
para("Learner Lab ở trạng thái Ready trong lần kiểm tra ngày 07/10/2026. Trước khi triển khai vào ngày 08/10, lệnh sts get-caller-identity trả về tài khoản 213858341631. CloudFormation chấp nhận template; stack learnerlab-ops-week1 đạt CREATE_COMPLETE lúc 12:52:18 (giờ Việt Nam). Lệnh ListStackResources cho thấy cả 10 tài nguyên đều hoàn tất. Mã băm SHA-256 của template và lệnh triển khai được lưu trong biên bản AWS kèm theo.")
table(["Thành phần", "Thông tin trên AWS"], [
    ("Account / Region", "213858341631 / us-east-1"),
    ("VPC", "vpc-0206a78d96a4b3a57; CIDR 10.20.0.0/16"),
    ("Subnet A", "subnet-0fe1f68d4d015130c; us-east-1a; 10.20.1.0/24"),
    ("Subnet B", "subnet-0a26fc106be139906; us-east-1b; 10.20.2.0/24"),
    ("Route table / IGW", "rtb-0d686db1acc56489d / igw-0998040b443aaa624"),
    ("S3 bucket", "learnerlab-ops-week1-artifactbucket-u4t7cql7ngfo"),
    ("IAM role có sẵn", "arn:aws:iam::213858341631:role/LabRole"),
], [4.3, 12.0])
caption("Bảng 3.1. Tài nguyên triển khai ngày 08/10/2026")

heading("3.2 Cấu hình bảo mật và tag", 2)
para("Bucket S3 bật cả bốn tùy chọn Block Public Access, mã hóa mặc định bằng SSE-S3 (AES256) và đặt Object Ownership là BucketOwnerEnforced. Các phần tải lên nhiều lượt nhưng chưa hoàn tất sẽ được dọn sau 7 ngày. Tài nguyên có hỗ trợ tag mang Project=learnerlab-ops, Owner=24110019 và Environment=learner-lab. Bucket dùng DeletionPolicy: Retain, nên khi xóa stack vẫn phải kiểm tra và xóa bucket riêng nếu không còn sử dụng.")
para("Lệnh IAM GetRole cho thấy LabRole đã tồn tại. Template có đầu ra LabInstanceProfileName để dùng về sau, nhưng chưa kiểm tra profile này bằng IAM API. Tuần 1 không tạo IAM user hoặc role mới.")

heading("Chương 4 Minh chứng và chi phí", 1)
heading("4.1 Ảnh chụp và nội dung kiểm tra", 2)
table(["Ảnh", "Nội dung thể hiện"], [
    ("01-lab-session-ready.jpg", "Trạng thái Ready, Account ID và Region ngày 07/10"),
    ("02-stack-complete.jpg", "Stack CREATE_COMPLETE, 10 tài nguyên và đầu ra VPC/subnet/S3"),
    ("03-network-check.jpg", "VPC, hai subnet, tuyến qua Internet Gateway và chế độ cấp public IPv4"),
    ("04-s3-private.jpg", "Block Public Access, mã hóa AES256, Object Ownership và tag"),
], [5.1, 11.2])
caption("Bảng 4.1. Ảnh AWS kèm Account ID, Region và thời điểm chụp")
para("Ảnh ngày 08/10 hiển thị giờ UTC−07:00, chậm hơn giờ Việt Nam 14 giờ. Ảnh và repository không chứa access key, secret key hoặc session token. Sau khi kiểm tra, tệp thông tin truy cập AWS CLI tạm thời được xóa và phiên Lab được kết thúc lúc khoảng 13:00 giờ Việt Nam.")

heading("4.2 Chi phí tại thời điểm kiểm tra", 2)
para("Cổng Learner Lab hiển thị Used $0 of $50 trong lần kiểm tra lại ngày 10/10/2026. Console EC2 không liệt kê instance ở cả us-east-1 và us-west-2. Kết thúc kiểm tra, phiên Lab hiển thị Terminated. AWS cho biết số dư ngân sách có thể cập nhật chậm 8–12 giờ, nên con số trên cổng chưa thể xem là hóa đơn cuối cùng. Stack tuần 1 không tạo EC2 hay NAT Gateway. Cost Explorer chưa được gọi trong tuần 1, vì vậy báo cáo chưa có dữ liệu để lập biểu đồ chi phí.")
para("Bucket S3 được giữ cho tuần 2. Dung lượng lưu trữ và các request sau thời điểm kiểm tra vẫn có thể phát sinh phí. Khi mở Lab lần sau, cần xem lại tài nguyên còn tồn tại và số dư trên cổng.")

heading("Chương 5 Tiến độ và kế hoạch tiếp theo", 1)
heading("5.1 Nhật ký công việc", 2)
para("AWS CloudTrail của tài khoản 213858341631 tại us-east-1 ghi nhận các sự kiện quản trị dưới đây. Thời gian là giờ Việt Nam (UTC+07:00) và thể hiện thời điểm sự kiện API được ghi nhận, không phải thời lượng học.")
table(["Thời điểm", "Hoạt động", "Tài liệu liên quan"], [
    ("08/10/2026 12:38:11", "CloudTrail ghi nhận ConsoleLogin cho phiên Learner Lab gắn với mã SV 24110019", "CloudTrail; bảng đối chiếu trong 08-cloudtrail-audit.md"),
    ("08/10/2026 12:47:24–12:47:29", "Tạo và thực thi Change Set cho stack learnerlab-ops-week1", "CloudTrail; bảng đối chiếu trong 08-cloudtrail-audit.md"),
    ("08/10/2026 12:47:32–12:47:34", "Tạo VPC, Internet Gateway và bucket S3; đặt cấu hình mã hóa, chặn truy cập công khai, lifecycle", "CloudTrail; VPC và bucket trong biên bản AWS"),
    ("08/10/2026 12:47:43–12:47:58", "Cấu hình VPC, tạo subnet/route table, gắn Internet Gateway và các tuyến", "CloudTrail; bảng đối chiếu trong 08-cloudtrail-audit.md"),
    ("08/10/2026 12:52:18", "Stack learnerlab-ops-week1 đạt CREATE_COMPLETE; 10/10 tài nguyên hoàn tất", "Biên bản CloudFormation; ảnh 02"),
], [4.2, 8.8, 3.3])
caption("Bảng 5.1. Các mốc hoạt động AWS được đối chiếu bằng nhật ký")
para("Lê Minh Huy xác nhận đã đọc tài liệu, chỉnh sửa sơ đồ/template, chạy một số lệnh AWS và kiểm tra ảnh. Biên bản triển khai ghi LLM/AI (ChatGPT) trực tiếp chạy template CloudFormation và các lệnh xác minh được lưu tại đó. CloudTrail gắn sự kiện với phiên AWS của mã SV 24110019 nhưng không cho biết ai đã gõ từng lệnh. Vì vậy, báo cáo không quy các mốc API cho riêng sinh viên hoặc công cụ; cột giờ cá nhân trong worklog để trống do không có dữ liệu đo thời lượng học.")
heading("5.2 Sử dụng LLM/AI (ChatGPT)", 2)
para("Trong tuần 1, LLM/AI (ChatGPT) được dùng để tra cứu tài liệu AWS, gợi ý kiến trúc và tạo bản thảo sơ đồ, template, script, báo cáo. Sinh viên đã chỉnh sửa sơ đồ/template và kiểm tra kết quả. Ở phiên Learner Lab ngày 08/10/2026, công cụ trực tiếp chạy template CloudFormation, các lệnh kiểm tra và lưu ảnh CLI. Các thao tác này được nêu riêng để phân biệt với phần sinh viên tự thực hiện. Sinh viên chịu trách nhiệm giải thích cấu hình, mã nguồn và kết quả khi trình bày báo cáo.")
heading("5.3 Kết quả tuần 1 và kế hoạch tuần 2", 2)
para("Sản phẩm tuần 1 gồm tài liệu yêu cầu, sơ đồ kiến trúc, stack VPC/S3, thông tin LabRole, tag Project/Owner/Environment và ảnh chụp có Account ID, Region, thời gian. Worklog cần bổ sung thời lượng thực tế theo từng buổi khi có dữ liệu đối chiếu; LabInstanceProfile cũng cần được kiểm tra nếu sẽ dùng ở giai đoạn sau.")
para("Trong tuần 2, đề tài sẽ triển khai Lambda, EventBridge, DynamoDB, API và các chức năng chính của dashboard. Trước khi đánh giá kết quả, cần kiểm tra quyền truy cập và dữ liệu thực tế của từng dịch vụ trong Learner Lab.")
heading("Nguồn và tệp kèm theo", 1)
source_p = para("Nguồn dữ kiện AWS: docs/week-1/07-deployment-record.md; docs/week-1/04-evidence-and-review.md; evidence/week-1/01–04. Nguồn kiến trúc: docs/week-1/01-services-and-requirements.md; docs/week-1/02-architecture.md; docs/diagrams/architecture-week1-simple.png. Mã hạ tầng: infra/week1-foundation.yaml. Repository GitHub công khai: ")
hyperlink(source_p, "github.com/leminhhuy-ute/learnerlab-ops-dashboard",
          "https://github.com/leminhhuy-ute/learnerlab-ops-dashboard")

heading("Phụ lục Ảnh minh chứng AWS", 1)
picture(ROOT / "evidence" / "week-1" / "02-stack-complete.jpg", 14.2,
        "Phụ lục A. Stack CREATE_COMPLETE, Account 213858341631, us-east-1")
picture(ROOT / "evidence" / "week-1" / "03-network-check.jpg", 14.2,
        "Phụ lục B. VPC, subnet và tuyến Internet Gateway")
picture(ROOT / "evidence" / "week-1" / "04-s3-private.jpg", 14.2,
        "Phụ lục C. Bảo mật bucket S3 và tag Project/Owner")

doc.core_properties.title = "Báo cáo tiến độ tuần 1 Bảng điều khiển vận hành cho AWS Learner Lab"
doc.core_properties.author = "Lê Minh Huy"
doc.core_properties.subject = "Báo cáo tuần 1 đề tài Monitoring & Operations"
OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
